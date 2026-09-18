from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .models import Order


def export_orders_xlsx(orders):
    """Таблица заказов «как на экране» (ТЗ): номер, покупатель, сумма,
    статус, дата — без агрегации по продажам (это отдельный отчёт, apps.reports)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заказы"

    sheet.append(["№ заказа", "Покупатель", "E-mail", "Сумма", "Статус", "Дата"])
    for order in orders:
        sheet.append([
            order.id,
            order.user.name,
            order.user.email,
            float(order.total_amount),
            Order.Status(order.status).label,
            order.created_at.strftime("%Y-%m-%d %H:%M"),
        ])

    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(length + 2, 60)

    buffer = BytesIO()
    workbook.save(buffer)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="orders.xlsx"'
    return response
