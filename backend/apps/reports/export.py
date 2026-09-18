import csv
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def _filename(report, extension):
    return f"sales-report-{report['date_from']}-{report['date_to']}.{extension}"


def export_csv(report):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "csv")}"'
    response.write("﻿")  # BOM — иначе Excel на Windows ломает кириллицу в CSV
    writer = csv.writer(response)
    writer.writerow(["Период", f"{report['date_from']} — {report['date_to']}"])
    writer.writerow(["Количество заказов", report["order_count"]])
    writer.writerow(["Общая выручка", report["total_revenue"]])
    writer.writerow([])
    writer.writerow(["Топ-10 продаваемых книг"])
    writer.writerow(["Книга", "Продано, шт.", "Выручка"])
    for book in report["top_books"]:
        writer.writerow([book["title"], book["quantity_sold"], book["revenue"]])
    return response


def export_xlsx(report):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт о продажах"

    sheet.append(["Период", f"{report['date_from']} — {report['date_to']}"])
    sheet.append(["Количество заказов", report["order_count"]])
    sheet.append(["Общая выручка", float(report["total_revenue"])])
    sheet.append([])
    sheet.append(["Топ-10 продаваемых книг"])
    sheet.append(["Книга", "Продано, шт.", "Выручка"])
    for book in report["top_books"]:
        sheet.append([book["title"], book["quantity_sold"], float(book["revenue"])])

    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(length + 2, 60)

    buffer = BytesIO()
    workbook.save(buffer)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "xlsx")}"'
    return response
